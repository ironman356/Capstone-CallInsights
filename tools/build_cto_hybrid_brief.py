"""Build a one-page CTO discussion handout for the Call Insights project."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("output/docs/Call_Insights_CTO_One_Pager.docx")
NAVY, BLUE, TEAL, INK, MUTED = "17324D", "2E74B5", "237B7B", "18212B", "536170"
PALE_BLUE, PALE_TEAL, PALE_GOLD = "EEF5FB", "EAF5F3", "FFF7E6"


def font(run, size, color=INK, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def border(paragraph, color=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), color)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def para(doc, text="", *, size=9.5, color=INK, after=4, before=0, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    font(p.add_run(text), size, color, bold, italic)
    return p


def label_para(doc, label, text, *, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.02
    font(p.add_run(label + " "), 9.3, NAVY, True)
    font(p.add_run(text), 9.3)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), 10.6, NAVY, True)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Inches(0.48), Inches(0.46)
    section.left_margin, section.right_margin = Inches(0.62), Inches(0.62)
    section.header_distance, section.footer_distance = Inches(0.2), Inches(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    font(header.add_run("CALL INSIGHTS  /  CTO DISCUSSION"), 8, MUTED, True)

    para(doc, "ONE-PAGE PROJECT UPDATE", size=8.5, color=BLUE, bold=True, after=2)
    para(doc, "Finding the Agent Approaches That Work", size=22, color=NAVY, bold=True, after=2)
    para(doc, "Simple goal: for the same customer problem, show which agent approaches lead to better outcomes and why.", size=10.4, color=MUTED, after=7)

    callout = para(doc, "We are not trying to make another dashboard. We are building the analysis underneath it: a way to spot useful explanations, phrases, and analogies that are currently buried across thousands of calls.", size=10, after=7)
    border(callout, TEAL)
    shade(callout, PALE_TEAL)

    heading(doc, "What the system does")
    flow = para(doc, "CALL TRANSCRIPT   ->   SEPARATE CUSTOMER PROBLEMS   ->   FIND THE AGENT'S APPROACH   ->   CHECK THE OUTCOME   ->   COMPARE SIMILAR CASES", size=8.2, color=NAVY, bold=True, after=4)
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(flow, PALE_BLUE)
    para(doc, "Example: two customers may both call about escrow. One agent gives a step-by-step explanation, while another only repeats a policy. We compare those approaches within the same problem instead of mixing unrelated calls.", size=9.2, after=3)

    heading(doc, "What we tested")
    label_para(doc, "Traditional ML:", "rules, TF-IDF, and older semantic models. They are fast and useful for a baseline, but wording changes and multi-problem calls are hard for them.")
    label_para(doc, "Hybrid LLM + ML:", "the LLM reads the conversation and finds the separate problems; ML retrieves approved topic options; the LLM chooses from those options and must point to the customer turns that support the decision.")
    label_para(doc, "Safety checks:", "if the system is unsure or sees a new type of issue, it says so instead of forcing a bad label. New clusters stay in review until a person approves them.")

    heading(doc, "Why the hybrid approach is the direction we recommend")
    benefits = para(doc, "CONTEXT  +  CONSISTENCY  +  PROOF", size=9.5, color=NAVY, bold=True, after=3)
    benefits.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(benefits, PALE_GOLD)
    para(doc, "The LLM understands different wording, analogies, and calls with more than one issue. The ML layer keeps reporting tied to stable business categories. Evidence-turn checks let a manager open the exact part of the call behind a result.", size=9.2, after=3)

    heading(doc, "Early results - promising, but not an SPS claim")
    para(doc, "On a blind synthetic test, the hybrid workflow found 46 issue segments across 28 calls and matched the known topic sequence correctly. The best traditional baseline reached 0.83 macro-F1 when it was already given the correct issue boundaries.", size=9.2, after=2)
    caveat = para(doc, "Important: this proves the workflow on our test data, not that it is production-ready. The next real test is 50-100 anonymized SPS calls labeled by reviewers.", size=8.8, color="7A5A00", bold=True, after=3)
    border(caveat, "C78C1B")

    heading(doc, "What this can help management find")
    para(doc, "- A phrase or analogy used by only a few agents that gets a better result\n- Which approaches work across many agents versus only for certain agents\n- What makes an approach effective: timing, clarity, empathy, a specific explanation, or a next step\n- Real examples with counts and uncertainty, not just a black-box score", size=9.0, after=4)

    heading(doc, "Decisions and help we need from SPS")
    para(doc, "1. A small anonymized, human-labeled call sample for evaluation.   2. Guidance on the approved GPU/server setup and local model options.   3. Confirmation of any internal practices or language that should not appear in a demo.", size=9.2, after=0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("Call Insights Capstone Team"), 8, MUTED)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
